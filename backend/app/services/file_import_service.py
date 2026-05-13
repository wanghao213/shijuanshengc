"""文件导入服务 - 支持多种格式文档的解析和处理."""

import asyncio
import tempfile
from pathlib import Path

import structlog
from sqlalchemy.ext.asyncio import AsyncSession

logger = structlog.get_logger()


class FileImportService:
    """文件导入服务，支持多种格式的文档解析."""

    def __init__(self, session: AsyncSession):
        self.session = session

    async def parse_docx(self, file_path: str) -> str | None:
        """解析 Word (.docx) 文件，提取文本内容."""
        try:
            from docx import Document
            
            doc = Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # 也提取表格内容
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text for cell in row.cells)
                    if row_text.strip():
                        paragraphs.append(row_text)
            
            text = "\n\n".join(paragraphs)
            logger.info("docx_parsed", path=file_path, length=len(text))
            return text
            
        except ImportError:
            logger.warning("python_docx_not_installed")
            return None
        except Exception as e:
            logger.error("docx_parse_error", error=str(e))
            return None

    async def parse_doc_legacy(self, file_path: str) -> str | None:
        """解析旧版 Word (.doc) 文件."""
        try:
            import subprocess
            
            # 尝试使用 antiword 或 catdoc
            for tool in ["antiword", "catdoc"]:
                try:
                    proc = await asyncio.create_subprocess_exec(
                        tool, file_path,
                        stdout=asyncio.subprocess.PIPE,
                        stderr=asyncio.subprocess.PIPE,
                    )
                    stdout, _ = await asyncio.wait_for(proc.communicate(), timeout=60)
                    if proc.returncode == 0 and stdout:
                        text = stdout.decode("utf-8", errors="ignore")
                        logger.info("doc_parsed_with_tool", tool=tool, path=file_path)
                        return text
                except FileNotFoundError:
                    continue
            
            logger.warning("no_doc_parser_available")
            return None
            
        except Exception as e:
            logger.error("doc_parse_error", error=str(e))
            return None

    async def parse_pdf_text(self, file_path: str) -> str | None:
        """解析 PDF 文件（纯文本模式，不使用 OCR）."""
        try:
            import fitz  # PyMuPDF
            
            doc = fitz.open(file_path)
            texts = []
            
            for page in doc:
                text = page.get_text()
                if text.strip():
                    texts.append(text)
            
            doc.close()
            result = "\n\n".join(texts)
            logger.info("pdf_text_parsed", path=file_path, pages=len(texts))
            return result
            
        except ImportError:
            logger.warning("pymupdf_not_installed")
            return None
        except Exception as e:
            logger.error("pdf_parse_error", error=str(e))
            return None

    async def parse_excel(self, file_path: str) -> str | None:
        """解析 Excel 文件，转换为文本格式."""
        try:
            import pandas as pd
            
            # 读取所有 sheet
            xls = pd.ExcelFile(file_path)
            sheets_text = []
            
            for sheet_name in xls.sheet_names:
                df = pd.read_excel(xls, sheet_name=sheet_name)
                # 转换为 Markdown 表格格式
                markdown_table = df.to_markdown(index=False)
                sheets_text.append(f"## {sheet_name}\n\n{markdown_table}")
            
            result = "\n\n".join(sheets_text)
            logger.info("excel_parsed", path=file_path, sheets=len(sheets_text))
            return result
            
        except ImportError:
            logger.warning("pandas_not_installed")
            return None
        except Exception as e:
            logger.error("excel_parse_error", error=str(e))
            return None

    async def get_file_info(self, file_path: str, filename: str) -> dict:
        """获取文件基本信息."""
        path = Path(file_path)
        suffix = path.suffix.lower()
        
        info = {
            "filename": filename,
            "suffix": suffix,
            "size_bytes": path.stat().st_size,
            "type_category": self._categorize_file(suffix),
        }
        
        # 根据文件类型提供额外信息
        if suffix in ['.pdf', '.doc', '.docx']:
            info["supports_ocr"] = True
        elif suffix in ['.xlsx', '.xls']:
            info["supports_structured_import"] = True
        elif suffix in ['.txt', '.md']:
            info["supports_direct_text"] = True
        
        return info

    def _categorize_file(self, suffix: str) -> str:
        """将文件按类型分类."""
        categories = {
            'document': ['.pdf', '.doc', '.docx', '.txt', '.md'],
            'spreadsheet': ['.xlsx', '.xls', '.csv'],
            'image': ['.png', '.jpg', '.jpeg', '.bmp', '.tiff'],
            'archive': ['.zip', '.rar'],
        }
        
        for category, suffixes in categories.items():
            if suffix in suffixes:
                return category
        
        return 'unknown'

    async def validate_file(self, file_path: str, max_size_mb: int = 50) -> tuple[bool, str | None]:
        """验证文件是否合法."""
        path = Path(file_path)
        
        if not path.exists():
            return False, "文件不存在"
        
        size_mb = path.stat().st_size / (1024 * 1024)
        if size_mb > max_size_mb:
            return False, f"文件大小超过限制 ({max_size_mb}MB)"
        
        suffix = path.suffix.lower()
        allowed_suffixes = {
            '.pdf', '.doc', '.docx', '.txt', '.md',
            '.xlsx', '.xls', '.csv',
            '.png', '.jpg', '.jpeg', '.bmp', '.tiff'
        }
        
        if suffix not in allowed_suffixes:
            return False, f"不支持的文件格式：{suffix}"
        
        return True, None
